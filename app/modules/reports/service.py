from pathlib import Path

from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.modules.billing.models import SalesInvoice
from app.modules.customers.models import CustomerLedgerEntry
from app.modules.products.models import Product


class ReportService:
    def __init__(self, db):
        self.db = db

    def export_sales_excel(self, output: str | Path):
        invoices = self.db.query(SalesInvoice).order_by(SalesInvoice.invoice_date).all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Invoice", "Date", "Taxable", "Tax", "Total", "Status"])
        for invoice in invoices:
            sheet.append([
                invoice.invoice_number, invoice.invoice_date, float(invoice.taxable_total),
                float(invoice.tax_total), float(invoice.total),
                "CANCELLED" if invoice.cancelled_at else "POSTED",
            ])
        workbook.save(output)
        return Path(output)

    def export_stock_excel(self, output: str | Path):
        products = self.db.query(Product).order_by(Product.name).all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["SKU", "Product", "Stock", "Unit Cost", "Stock Value"])
        for product in products:
            sheet.append([
                product.sku, product.name, product.current_stock, float(product.unit_cost),
                float(product.unit_cost) * product.current_stock,
            ])
        workbook.save(output)
        return Path(output)

    def export_udhaar_excel(self, customer_id: str, output: str | Path):
        entries = self.db.query(CustomerLedgerEntry).filter(
            CustomerLedgerEntry.customer_id == customer_id
        ).order_by(CustomerLedgerEntry.entry_date).all()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Date", "Type", "Amount", "Signed Amount", "Description"])
        for entry in entries:
            sheet.append([
                entry.entry_date, entry.entry_type.value, float(entry.amount),
                float(entry.signed_amount), entry.description,
            ])
        workbook.save(output)
        return Path(output)

    def export_invoice_pdf(self, invoice_id: str, output: str | Path):
        invoice = self.db.get(SalesInvoice, invoice_id)
        if invoice is None:
            raise ValueError("Invoice not found.")
        pdf = canvas.Canvas(str(output))
        pdf.drawString(50, 800, f"Invoice {invoice.invoice_number}")
        pdf.drawString(50, 780, f"Date: {invoice.invoice_date}")
        y = 750
        for line in invoice.lines:
            pdf.drawString(50, y, f"{line.description} x {line.quantity}: {line.total:.2f}")
            y -= 20
        pdf.drawString(50, y - 10, f"Total: {invoice.total:.2f}")
        pdf.save()
        return Path(output)
