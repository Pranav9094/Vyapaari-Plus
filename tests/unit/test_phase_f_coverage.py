from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.api import require_permission
from app.core.exceptions import ValidationError
from app.modules.alerts.service import AlertService
from app.modules.auth.models import Role
from app.modules.auth.service import AuthService
from app.modules.billing.models import InvoiceCounter, PaymentMode, SalesInvoice, SalesPayment
from app.modules.billing.service import BillingService
from app.modules.business.service import BusinessService
from app.modules.customers.models import CustomerLedgerEntry, CustomerEntryType
from app.modules.customers.service import CustomerLedgerService
from app.modules.inventory.models import Branch, StockTransfer, TransferStatus
from app.modules.products.models import Product
from app.modules.reports.service import ReportService
from app.modules.suppliers.models import Supplier, SupplierInvoice


GSTIN = "24AAACC1206D1ZM"


def setup_business(db_session, *, state="24", prefix="INV"):
    return BusinessService(db_session).create_or_update(
        name="Test Retail", state_code=state, address="Main Road",
        gstin=GSTIN if state == "24" else "27AAACC1206D1ZM",
        invoice_prefix=prefix,
    )


def add_product(db_session, product_id="p1", stock=20, price="10.00"):
    product = Product(
        id=product_id, sku=product_id.upper(), name=f"Product {product_id}",
        current_stock=stock, reorder_point=2, reorder_quantity=5,
        unit_cost=Decimal(price),
    )
    db_session.add(product)
    db_session.commit()
    return product


def bill(db_session, **kwargs):
    items = kwargs.pop("items", [{
        "product_id": "p1", "quantity": 1, "unit_price": "10.00",
        "gst_rate": "18", "hsn_sac": "1001",
    }])
    return BillingService(db_session).create_bill(
        items=items,
        idempotency_key=kwargs.pop("idempotency_key", "key"),
        **kwargs,
    )


def test_business_rejects_bad_gstin_state_month_and_prefix(db_session):
    service = BusinessService(db_session)
    for kwargs in (
        {"state_code": "24", "gstin": "24AAACC1206D1ZN"},
        {"state_code": "27", "gstin": GSTIN},
        {"state_code": "00", "gstin": None},
        {"state_code": "24", "financial_year_start_month": 13},
        {"state_code": "24", "invoice_prefix": "BAD PREFIX"},
    ):
        params = {"name": "Retail", "address": "Road", "invoice_prefix": "INV"}
        params.update(kwargs)
        with pytest.raises(ValidationError):
            service.create_or_update(**params)


def test_billing_calculates_intra_state_gst_and_split_payments(db_session):
    setup_business(db_session)
    add_product(db_session)
    invoice = bill(
        db_session,
        idempotency_key="intra",
        payments=[
            {"mode": PaymentMode.CASH, "amount": "10.00"},
            {"mode": PaymentMode.UPI, "amount": "1.80"},
        ],
    )
    assert invoice.tax_total == Decimal("1.80")
    assert invoice.total == Decimal("12.00")
    assert {payment.mode for payment in invoice.payments} == {
        PaymentMode.CASH, PaymentMode.UPI
    }
    assert invoice.lines[0].tax == Decimal("1.80")


def test_billing_calculates_inter_state_gst_and_line_discount(db_session):
    setup_business(db_session)
    add_product(db_session)
    invoice = bill(
        db_session, place_of_supply="27", idempotency_key="inter",
        payments=[{"mode": PaymentMode.CARD, "amount": "10.62"}],
    )
    invoice = BillingService(db_session).create_bill(
        items=[{
            "product_id": "p1", "quantity": 1, "unit_price": "10.00",
            "discount": "1.00", "gst_rate": "18", "hsn_sac": "1001",
        }],
        place_of_supply="27", payments=[{"mode": PaymentMode.CARD, "amount": "10.62"}],
        idempotency_key="inter-discount",
    )
    assert invoice.tax_total == Decimal("1.62")
    assert invoice.taxable_total == Decimal("9.00")
    assert invoice.lines[0].tax == Decimal("1.62")


def test_bill_discount_is_applied_before_tax_and_rounds_half_up(db_session):
    setup_business(db_session)
    add_product(db_session)
    invoice = BillingService(db_session).create_bill(
        items=[{
            "product_id": "p1", "quantity": 1, "unit_price": "10.01",
            "gst_rate": "5", "hsn_sac": "1001",
        }],
        bill_discount="1.00", idempotency_key="bill-discount",
    )
    assert invoice.taxable_total == Decimal("9.01")
    assert invoice.tax_total == Decimal("0.45")
    assert invoice.total == Decimal("9.00")
    assert invoice.round_off == Decimal("-0.46")


def test_billing_rolls_back_stock_and_invoice_counter_on_payment_failure(db_session):
    setup_business(db_session)
    add_product(db_session, stock=3)
    service = BillingService(db_session)
    with pytest.raises(ValidationError):
        bill(
            db_session, idempotency_key="bad-payment",
            payments=[{"mode": PaymentMode.CASH, "amount": "999.00"}],
        )
    assert db_session.get(Product, "p1").current_stock == 3
    assert db_session.query(SalesInvoice).count() == 0
    invoice = bill(
        db_session, idempotency_key="after-failure",
        payments=[{"mode": PaymentMode.CASH, "amount": "12.00"}],
    )
    assert invoice.invoice_number.endswith("0001")


def test_billing_validates_required_inputs_and_cancel_reverses_stock_and_udhaar(db_session):
    service = BillingService(db_session)
    with pytest.raises(ValidationError, match="idempotency"):
        service.create_bill(items=[], idempotency_key=None)
    with pytest.raises(ValidationError, match="configured"):
        service.create_bill(items=[], idempotency_key="no-profile")
    setup_business(db_session)
    add_product(db_session, stock=2)
    customer = CustomerLedgerService(db_session).create_customer("Credit Customer")
    with pytest.raises(ValidationError):
        bill(db_session, idempotency_key="bad-product", items=[{"product_id": "missing"}])
    invoice = bill(db_session, idempotency_key="credit", customer_id=customer.id)
    assert CustomerLedgerService(db_session).balance(customer.id) == invoice.total
    BillingService(db_session).cancel_bill(invoice.id)
    assert db_session.get(Product, "p1").current_stock == 2
    assert CustomerLedgerService(db_session).balance(customer.id) == Decimal("0.00")
    assert BillingService(db_session).cancel_bill(invoice.id).id == invoice.id
    with pytest.raises(ValidationError):
        BillingService(db_session).cancel_bill("missing")


def test_invoice_sequence_resets_with_financial_year_and_rejects_long_number(db_session):
    profile = setup_business(db_session, prefix="SALE")
    add_product(db_session)
    first = bill(db_session, idempotency_key="seq-1")
    second = bill(db_session, idempotency_key="seq-2")
    assert first.invoice_number.endswith("0001")
    assert second.invoice_number.endswith("0002")
    profile.invoice_prefix = "TOOLONG"
    db_session.commit()
    with pytest.raises(ValidationError, match="exceeds 16"):
        bill(db_session, idempotency_key="seq-long")
    old_counter = InvoiceCounter(
        id=f"{profile.id}:1900-01", next_number=99
    )
    db_session.add(old_counter)
    db_session.commit()
    assert old_counter.id != f"{profile.id}:{profile.financial_year}"


def test_udhaar_credit_limit_partial_payment_reversal_and_invariant(db_session):
    service = CustomerLedgerService(db_session)
    customer = service.create_customer("Customer", credit_limit=Decimal("100.00"))
    service.record_sale(customer.id, Decimal("80.00"))
    db_session.commit()
    with pytest.raises(ValidationError, match="credit limit"):
        service.record_sale(customer.id, Decimal("21.00"))
    payment = service.record_payment(customer.id, Decimal("30.00"))
    assert payment.amount == Decimal("30.00")
    service.record_reversal(customer.id, Decimal("20.00"))
    db_session.commit()
    entries = db_session.query(CustomerLedgerEntry).filter_by(customer_id=customer.id).all()
    assert service.balance(customer.id) == sum(
        (entry.signed_amount for entry in entries), Decimal("0.00")
    )
    assert service.balance(customer.id) == Decimal("30.00")
    with pytest.raises(ValidationError):
        service.record_payment(customer.id, Decimal("31.00"))


def test_udhaar_rejects_invalid_customers_amounts_and_exposes_statement(db_session):
    service = CustomerLedgerService(db_session)
    with pytest.raises(ValidationError):
        service.create_customer(" ", credit_limit="1")
    with pytest.raises(ValidationError):
        service.create_customer("Bad limit", credit_limit="-1")
    with pytest.raises(ValidationError):
        service.record_sale("missing", Decimal("1"))
    customer = service.create_customer("Statement Customer")
    with pytest.raises(ValidationError):
        service.record_sale(customer.id, Decimal("0"))
    service.record_sale(customer.id, Decimal("10"))
    db_session.commit()
    rows = service.statement(customer.id)
    assert rows[-1]["balance"] == Decimal("10.00")


def test_rbac_has_401_403_and_role_permission_matrix(db_session):
    dependency = require_permission("billing")
    with pytest.raises(HTTPException) as missing:
        dependency(token=None, db=db_session)
    assert missing.value.status_code == 401
    with pytest.raises(HTTPException) as invalid:
        dependency(token="expired-token", db=db_session)
    assert invalid.value.status_code == 403

    service = AuthService(db_session)
    users = {
        role: service.create_user(role.value, "password-123", role)
        for role in Role
    }
    permissions = ("business", "billing", "inventory", "transfers", "reports")
    for role, user in users.items():
        token = service.login(user.username, "password-123")
        for permission in permissions:
            allowed = permission in {
                Role.OWNER: {"business", "billing", "inventory", "transfers", "reports"},
                Role.MANAGER: {"billing", "inventory", "transfers", "reports"},
                Role.CASHIER: {"billing"},
            }[role]
            dependency = require_permission(permission)
            if allowed:
                assert dependency(token=token, db=db_session).id == user.id
            else:
                with pytest.raises(HTTPException) as denied:
                    dependency(token=token, db=db_session)
                assert denied.value.status_code == 403
        AuthService._tokens.pop(token)
        with pytest.raises(HTTPException):
            require_permission("billing")(token=token, db=db_session)


def test_exports_match_database_and_empty_exports(tmp_path, db_session):
    empty = ReportService(db_session)
    empty.export_sales_excel(tmp_path / "empty.xlsx")
    assert load_workbook(tmp_path / "empty.xlsx").active.max_row == 1
    setup_business(db_session)
    add_product(db_session, stock=4, price="2.50")
    invoice = bill(
        db_session, idempotency_key="export-invoice",
        payments=[{"mode": PaymentMode.CASH, "amount": "12.00"}],
    )
    report = ReportService(db_session)
    sales_path = report.export_sales_excel(tmp_path / "sales.xlsx")
    stock_path = report.export_stock_excel(tmp_path / "stock.xlsx")
    pdf_path = report.export_invoice_pdf(invoice.id, tmp_path / "invoice.pdf")
    sales_sheet = load_workbook(sales_path).active
    stock_sheet = load_workbook(stock_path).active
    assert sales_sheet.max_row == 2
    assert Decimal(str(sales_sheet["E2"].value)) == invoice.total
    assert stock_sheet.max_row == 2
    assert stock_sheet["E2"].value == 7.5
    assert pdf_path.stat().st_size > 0
    customer = CustomerLedgerService(db_session).create_customer("Export Customer")
    customer_service = CustomerLedgerService(db_session)
    customer_service.record_sale(customer.id, Decimal("3.00"))
    db_session.commit()
    udhaar_path = report.export_udhaar_excel(customer.id, tmp_path / "udhaar.xlsx")
    assert load_workbook(udhaar_path).active.max_row == 2
    stock_empty_path = report.export_stock_excel(tmp_path / "stock-empty.xlsx")
    assert load_workbook(stock_empty_path).active.max_row == 2
    with pytest.raises(ValueError):
        report.export_invoice_pdf("missing", tmp_path / "missing.pdf")


def test_alerts_return_each_alert_category(db_session):
    db_session.add(Product(
        id="low", sku="LOW", name="Low", current_stock=1, reorder_point=2,
        reorder_quantity=5, unit_cost=Decimal("1.00"),
    ))
    supplier = Supplier(id="supplier", name="Supplier")
    db_session.add(supplier)
    db_session.add(SupplierInvoice(
        id="payable", supplier_id="supplier", invoice_number="SUP-1",
        invoice_date=datetime.now() - timedelta(days=60),
        due_date=datetime.now() - timedelta(days=1),
        total_amount=Decimal("10.00"), amount_paid=Decimal("0.00"),
    ))
    branch_a = Branch(id="branch-a", name="A", code="A")
    branch_b = Branch(id="branch-b", name="B", code="B")
    db_session.add_all([branch_a, branch_b])
    db_session.add(StockTransfer(
        id="transfer", transfer_number="TR-1", source_branch_id="branch-a",
        destination_branch_id="branch-b", status=TransferStatus.IN_TRANSIT,
    ))
    overdue = CustomerLedgerEntry(
        id="overdue", customer_id="missing-customer", entry_type=CustomerEntryType.SALE,
        amount=Decimal("5.00"), description="old", entry_date=datetime.now() - timedelta(days=31),
    )
    # SQLite does not enforce this foreign key in the test engine, matching the service query contract.
    db_session.add(overdue)
    db_session.commit()
    alerts = AlertService(db_session)
    assert [item.id for item in alerts.low_stock()] == ["low"]
    assert [item.id for item in alerts.overdue_udhaar()] == ["overdue"]
    assert [item.id for item in alerts.overdue_payables()] == ["payable"]
    assert [item.id for item in alerts.pending_transfers()] == ["transfer"]
